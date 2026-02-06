"""
deckAIn - Citations Document Builder
Stage 4: Create citations.docx with source traceability
"""

from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from extractors.schemas import FinancialData, PublicMarketData, Citation
from config.settings import OUTPUT_DIR, ENABLE_CITATIONS
from utils.logger import setup_logger, log_stage

logger = setup_logger(__name__)

class CitationsBuilder:
 """
 Build citations.docx documenting all data sources.
 Required for compliance and traceability.
 """
 
 def __init__(self):
     """Initialize citations builder."""
     logger.info("CitationsBuilder initialized")
 
 def build_citations_document(
     self,
     financial_data: FinancialData,
     public_data: PublicMarketData,
     company_codename: str,
     output_path: str = None
     ) -> str:
     """
     Create citations document with all sources.
     
     Args:
     financial_data: Financial data with source metadata
     public_data: Public market data with URLs
     company_codename: Anonymized project name
     output_path: Optional output path
     
     Returns:
     Path to created citations.docx
     """
     log_stage(logger, 4, "Citations Document", "START")
     
     if not ENABLE_CITATIONS:
         logger.info(" Citations disabled in config")
         log_stage(logger, 4, "Citations Document", "SKIP")
         return None
     
     # Create document
     doc = Document()
     
     # Title
     title = doc.add_heading(f'{company_codename} - Source Citations', 0)
     title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
     
     # Introduction
     intro = doc.add_paragraph(
         "This document provides complete traceability for all data points and claims "
         "in the investment teaser presentation. Every number and market context reference "
         "is mapped to its original source."
     )
     intro_format = intro.runs[0]
     intro_format.font.size = Pt(10)
     intro_format.font.color.rgb = RGBColor(100, 100, 100)
     
     doc.add_paragraph() # Spacer
     
     # Section 1: Private Financial Data Sources
     doc.add_heading('1. Private Financial Data Sources', level=1)
     self._add_private_data_citations(doc, financial_data)
     
     # Section 2: Public Market Context Sources
     doc.add_heading('2. Public Market Context Sources', level=1)
     self._add_public_data_citations(doc, public_data)
     
     # Section 3: Methodology
     doc.add_heading('3. Data Processing Methodology', level=1)
     self._add_methodology_section(doc, financial_data)
     
     # Save document
     if output_path is None:
         output_path = OUTPUT_DIR / f"{company_codename.replace(' ', '_')}_citations.docx"
     else:
         output_path = Path(output_path)
     
     output_path.parent.mkdir(parents=True, exist_ok=True)
     doc.save(str(output_path))
     
     log_stage(logger, 4, "Citations Document", "DONE")
     logger.info(f" Citations saved: {output_path}")
     
     return str(output_path)
 
 def _add_private_data_citations(self, doc: Document, data: FinancialData):
     """Add private data source citations."""
     
     # Source file information
     if data.source_metadata:
         doc.add_heading('Source Document', level=2)
         p = doc.add_paragraph()
         p.add_run('File: ').bold = True
         p.add_run(f"{data.source_metadata.source_file}\n")
         p.add_run('Extraction Date: ').bold = True
         p.add_run(f"{data.source_metadata.extraction_timestamp}\n")
         p.add_run('Primary Table: ').bold = True
         p.add_run(f"{data.source_metadata.table_name}")
     
     # Financial metrics
     doc.add_heading('Financial Metrics Extracted', level=2)
     
     # Revenue
     if data.revenue_by_year and data.years:
         doc.add_heading('Revenue', level=3)
         for year, value in zip(data.years, data.revenue_by_year):
             p = doc.add_paragraph(f"{year}: ₹{value} Crores", style='List Bullet')
             p.runs[0].font.size = Pt(10)
     
     # EBITDA
     if data.ebitda_by_year and data.years:
         doc.add_heading('EBITDA', level=3)
         for year, value in zip(data.years, data.ebitda_by_year):
             p = doc.add_paragraph(f"{year}: ₹{value} Crores", style='List Bullet')
             p.runs[0].font.size = Pt(10)
     
     # PAT
     if data.pat_by_year and data.years:
         doc.add_heading('PAT (Profit After Tax)', level=3)
         for year, value in zip(data.years, data.pat_by_year):
             p = doc.add_paragraph(f"{year}: ₹{value} Crores", style='List Bullet')
             p.runs[0].font.size = Pt(10)
     
     # Other metrics
     doc.add_heading('Operational Metrics', level=3)
     
     if data.employee_count:
         p = doc.add_paragraph(f"Employee Count: {data.employee_count}", style='List Bullet')
         p.runs[0].font.size = Pt(10)
     
     if data.facility_count:
         p = doc.add_paragraph(f"Manufacturing Facilities: {data.facility_count}", style='List Bullet')
         p.runs[0].font.size = Pt(10)
     
     if data.export_revenue_pct:
         p = doc.add_paragraph(f"Export Revenue: {data.export_revenue_pct}%", style='List Bullet')
         p.runs[0].font.size = Pt(10)
 
 def _add_public_data_citations(self, doc: Document, public_data: PublicMarketData):
     """Add public data source citations."""
     
     all_sources = public_data.get_all_sources()
     
     if not all_sources:
         doc.add_paragraph("No public data sources used.")
         return
     
     doc.add_paragraph(
         f"Total public sources referenced: {len(all_sources)}"
     )
     
     # Group by category
     categories = {
         "Market Size & Opportunity": public_data.market_size_data,
         "Industry Trends": public_data.industry_trends,
         "Competitive Landscape": public_data.competitive_landscape,
         "Recent News": public_data.recent_news
     }
     
     for category, sources in categories.items():
         if sources:
             doc.add_heading(category, level=2)
             
             for idx, source in enumerate(sources, 1):
                 p = doc.add_paragraph()
                 p.add_run(f"Source {idx}: ").bold = True
                 p.add_run(f"{source.title}\n")
                 p.add_run("URL: ").bold = True
                 
                 # Add URL as plain text (hyperlink API is buggy)
                 p.add_run(f"{source.url}\n")
                 
                 p.add_run(f"Accessed: {source.timestamp}\n")
                 p.add_run("Excerpt: ").italic = True
                 p.add_run(f'"{source.snippet[:200]}..."')
                 
                 p.runs[0].font.size = Pt(9)
 
 def _add_methodology_section(self, doc: Document, data: FinancialData):
     """Add methodology and calculations section."""
     
     doc.add_paragraph(
         "All calculated metrics (CAGR, margins, ratios) were computed using Python "
         "mathematical functions based on the extracted raw data. The LLM was NOT used "
         "for calculations to ensure 100% accuracy."
     )
     
     # List calculations
     doc.add_heading('Calculated Metrics', level=2)
     
     if data.revenue_cagr:
         p = doc.add_paragraph()
         p.add_run(f"Revenue CAGR: {data.revenue_cagr}%\n").bold = True
         p.add_run("Formula: ((End Value / Start Value) ^ (1 / Years) - 1) × 100\n")
         p.add_run(f"Calculation: (({data.revenue_by_year[-1]} / {data.revenue_by_year[0]}) ^ ")
         p.add_run(f"(1 / {len(data.revenue_by_year) - 1}) - 1) × 100 = {data.revenue_cagr}%")
         p.runs[0].font.size = Pt(9)
     
     if data.ebitda_margin_latest:
         p = doc.add_paragraph()
         p.add_run(f"Latest EBITDA Margin: {data.ebitda_margin_latest}%\n").bold = True
         p.add_run("Formula: (EBITDA / Revenue) × 100\n")
         p.add_run(f"Calculation: ({data.ebitda_by_year[-1]} / {data.revenue_by_year[-1]}) × 100 = ")
         p.add_run(f"{data.ebitda_margin_latest}%")
         p.runs[0].font.size = Pt(9)

def add_hyperlink(paragraph, url, text):
 """
 Add a hyperlink to a paragraph.
 
 Args:
 paragraph: docx paragraph object
 url: URL to link to
 text: Display text
 """
 part = paragraph.part
 r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
 
 hyperlink = paragraph._element.add_hyperlink(r_id)
 run = hyperlink.add_r()
 run.text = text
 
 # Style the hyperlink
 run.rPr.rStyle.val = 'Hyperlink'
 run.rPr.color = RGBColor(0, 0, 255)
 run.rPr.u.val = 'single'
