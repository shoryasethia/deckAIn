"""
deckAIn - Visual Component Library
Reusable "Smart Components" for premium slide layouts.
"""

from pathlib import Path
from PIL import Image
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from config.settings import KELP_COLORS, KELP_FONT_HEAD, KELP_FONT_BODY
from utils.logger import setup_logger

logger = setup_logger(__name__)

def _hex_to_rgb(hex_color: str):
    """Convert hex string (e.g. '003366') to tuple (0, 51, 102)."""
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

def draw_section_header(slide, title: str, x: float, y: float, width: float = 9.0):
    """Draw a premium section header (underline style)."""
    # Title
    tbox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(0.5))
    p = tbox.text_frame.paragraphs[0]
    p.text = title.upper()
    p.font.name = KELP_FONT_HEAD
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["primary_dark"]))
    
    # Underline
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(x), Inches(y + 0.35), Inches(width), Inches(0.03)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["gradient_pink"]))
    line.line.fill.background()

def draw_timeline(slide, events: list, x: float, y: float, width: float, height: float):
    """
    Draw a horizontal timeline component.
    Args:
        events: List of dicts {'year': '2000', 'event': 'Text'}
    """
    # 1. Main horizontal numeric axis
    axis_y = y + (height / 2)
    axis = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(axis_y), Inches(width), Inches(0.05)
    )
    axis.fill.solid()
    axis.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["primary_indigo"]))
    axis.line.fill.background()
    
    # 2. Points
    if not events:
        return
        
    num_points = len(events)
    # spacing
    img_w = width
    step = img_w / (num_points + 1)
    
    for i, item in enumerate(events):
        cx = x + step * (i + 1)
        cy = axis_y
        
        # Dot
        dot_size = 0.2
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx - dot_size/2), Inches(cy - dot_size/2 + 0.025), # Center on line
            Inches(dot_size), Inches(dot_size)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["primary_indigo"]))
        dot.line.color.rgb = RGBColor(255, 255, 255)
        dot.line.width = Pt(2)
        
        # Year (Above)
        yb = slide.shapes.add_textbox(
            Inches(cx - 0.5), Inches(cy - 0.6), Inches(1.0), Inches(0.4)
        )
        yp = yb.text_frame.paragraphs[0]
        yp.text = item.get("year", "")
        yp.alignment = PP_ALIGN.CENTER
        yp.font.name = KELP_FONT_HEAD
        yp.font.size = Pt(12)
        yp.font.bold = True
        yp.font.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["primary_dark"]))
        
        # Description (Below)
        db = slide.shapes.add_textbox(
            Inches(cx - 0.75), Inches(cy + 0.2), Inches(1.5), Inches(1.0)
        )
        text_frame = db.text_frame
        text_frame.word_wrap = True
        dp = text_frame.paragraphs[0]
        dp.text = item.get("event", "")
        dp.alignment = PP_ALIGN.CENTER
        dp.font.name = KELP_FONT_BODY
        dp.font.size = Pt(9)
        dp.font.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["text_body"]))

def draw_stat_grid(slide, stats: list, x: float, y: float, width: float, height: float):
    """
    Draw a grid of KPI cards.
    Args:
        stats: List of dicts {'label': 'Rev', 'value': '$100M'}
    """
    if not stats:
        return
        
    # Layout: 2 items per row max
    import math
    max_cols = 2
    rows = math.ceil(len(stats) / max_cols)
    
    box_w = (width - 0.2) / max_cols # gap 0.2
    box_h = (height - 0.2 * (rows-1)) / rows
    
    for i, stat in enumerate(stats):
        r = i // max_cols
        c = i % max_cols
        
        bx = x + c * (box_w + 0.2)
        by = y + r * (box_h + 0.2)
        
        # Card Background (Rounded)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(bx), Inches(by), Inches(box_w), Inches(box_h)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(245, 245, 250) # Very light grey
        card.line.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["text_light"]))
        card.shadow.inherit = False # No default shadow mess
        
        # Value (Big)
        vb = slide.shapes.add_textbox(
            Inches(bx + 0.1), Inches(by + 0.08), Inches(box_w - 0.2), Inches(box_h * 0.4)
        )
        vp = vb.text_frame.paragraphs[0]
        value_text = str(stat.get("value", ""))
        vp.text = value_text
        vp.font.name = KELP_FONT_HEAD
        if len(value_text) <= 4:
            vp.font.size = Pt(16)
        elif len(value_text) <= 7:
            vp.font.size = Pt(14)
        else:
            vp.font.size = Pt(12)
        vp.font.bold = True
        vp.font.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["gradient_pink"]))
        vp.alignment = PP_ALIGN.CENTER
        
        # Label (Small)
        lb = slide.shapes.add_textbox(
            Inches(bx + 0.1), Inches(by + 0.45), Inches(box_w - 0.2), Inches(box_h * 0.4)
        )
        lp = lb.text_frame.paragraphs[0]
        lp.text = str(stat.get("label", ""))
        lp.font.name = KELP_FONT_BODY
        lp.font.size = Pt(8)
        lp.alignment = PP_ALIGN.CENTER
        lp.font.color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS["text_body"]))

# =============================================================================
# ROBUST PRIMITIVES (Crashes-Proof Wrappers for AI)
# =============================================================================

def draw_rect(slide, x, y, w, h, color_key="primary_indigo", rounded=False):
    """Safe rectangle drawing."""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    try:
        x, y, w, h = float(x), float(y), float(w), float(h)
    except:
        x, y, w, h = 0.5, 1.2, 1.0, 1.0 # Safe fallback
        
    rect = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    
    # Safe Fill
    try:
        rect.fill.solid()
        hex_val = KELP_COLORS.get(color_key, KELP_COLORS["primary_indigo"])
        rect.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(hex_val))
    except:
        pass
        
    # No border by default
    try:
        rect.line.fill.background()
    except:
        pass
    return rect

def draw_line(slide, x, y, x2, y2, color_key="primary_indigo", weight_pt=1.5):
    """Safe line drawing (fixes 'LineFormat' attribute errors)."""
    try:
        x, y, x2, y2 = float(x), float(y), float(x2), float(y2)
    except:
        x, y, x2, y2 = 0.5, 1.2, 1.5, 1.2
        
    # In pptx, 'LINE' is best created as a connector or a very thin rectangle for predictability
    from pptx.enum.shapes import MSO_CONNECTOR
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x2), Inches(y2))
    
    try:
        hex_val = KELP_COLORS.get(color_key, KELP_COLORS["primary_indigo"])
        # Correct way for lines: .line.color, not .fill.fore_color
        line.line.color.rgb = RGBColor(*_hex_to_rgb(hex_val))
        line.line.width = Pt(weight_pt)
    except:
        pass
    return line

def draw_text_box(slide, text, x, y, w, h, size=11, color_key="text_dark", bold=False, align=PP_ALIGN.LEFT):
    """Safe text box drawing."""
    try:
        x, y, w, h = float(x), float(y), float(w), float(h)
        size = float(size)
    except:
        x, y, w, h = 0.5, 1.2, 1.0, 1.0
        size = 11

    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = str(text)
    
    # Safe Alignment
    try:
        p.alignment = align
    except:
        logger.warning(f"Invalid alignment '{align}' requested. Defaulting to LEFT.")
        p.alignment = PP_ALIGN.LEFT
    
    try:
        p.font.name = KELP_FONT_BODY
        p.font.size = Pt(size)
        p.font.bold = bold
        hex_val = KELP_COLORS.get(color_key, KELP_COLORS["text_dark"])
        p.font.color.rgb = RGBColor(*_hex_to_rgb(hex_val))
    except:
        pass
    return tx

def draw_image_box(slide, image_path, x, y, w, h):
    """Safe image drawing with error handling (center-crop to fit box)."""
    if not image_path:
        return None
    try:
        x, y, w, h = float(x), float(y), float(w), float(h)
        image_path = _prepare_cover_image(image_path, w / h)
        return slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    except Exception as e:
        logger.debug(f"Failed to add image: {e}")
        return None

def _prepare_cover_image(image_path: str, target_ratio: float) -> str:
    """Center-crop image to match target aspect ratio, then save a temp copy."""
    try:
        path = Path(image_path)
        if not path.exists():
            return str(image_path)

        with Image.open(path) as img:
            width, height = img.size
            if height == 0:
                return str(path)

            current_ratio = width / height
            if abs(current_ratio - target_ratio) < 0.01:
                return str(path)

            if current_ratio > target_ratio:
                new_width = int(height * target_ratio)
                left = int((width - new_width) / 2)
                box = (left, 0, left + new_width, height)
            else:
                new_height = int(width / target_ratio)
                top = int((height - new_height) / 2)
                box = (0, top, width, top + new_height)

            cropped = img.crop(box)
            out_name = f"crop_{int(target_ratio * 1000)}_{path.name}"
            out_path = path.parent / out_name
            if not out_path.exists():
                cropped.save(out_path, quality=95, optimize=True)

            return str(out_path)
    except Exception as e:
        logger.debug(f"Image crop failed: {e}")
        return str(image_path)

def draw_bullet_point(slide, text, x, y, width=5.0, color_key="primary_indigo", size=11):
    """Draw a bullet point with colored bullet."""
    # Bullet circle
    bullet = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x), Inches(y + 0.05),
        Inches(0.15), Inches(0.15)
    )
    bullet.fill.solid()
    bullet.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS.get(color_key, KELP_COLORS["primary_indigo"])))
    bullet.line.fill.background()
    
    # Text
    text_box = draw_text_box(slide, text, x + 0.25, y, width - 0.25, 0.3, size=size, color_key="text_body")
    return text_box

def draw_decorative_bar(slide, x, y, width, color_key="gradient_orange", height=0.05):
    """Draw a decorative horizontal bar."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y),
        Inches(width), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(KELP_COLORS.get(color_key, KELP_COLORS["gradient_orange"])))
    bar.line.fill.background()
    return bar
