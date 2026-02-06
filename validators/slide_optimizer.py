"""
deckAIn - Slide Optimizer
Applies structured fixes from Groq review to PPTX slides.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from typing import Dict, List, Any
from pathlib import Path

from utils.logger import setup_logger

logger = setup_logger(__name__)


class SlideOptimizer:
    """
    Apply structured fixes to PPTX slides based on Groq feedback.
    Safe operations only - no arbitrary code execution.
    """
    
    def __init__(self):
        """Initialize optimizer."""
        logger.info("SlideOptimizer initialized")
    
    def apply_fixes(self, pptx_path: str, slide_index: int, fixes: List[Dict[str, Any]]) -> bool:
        """
        Apply a list of fixes to a specific slide.
        
        Args:
            pptx_path: Path to PPTX file
            slide_index: 0-indexed slide number
            fixes: List of fix dictionaries from Groq
            
        Returns:
            True if fixes applied successfully
        """
        try:
            prs = Presentation(pptx_path)
            
            if slide_index >= len(prs.slides):
                logger.error(f"Slide index {slide_index} out of range")
                return False
            
            slide = prs.slides[slide_index]
            applied_count = 0
            
            for fix in fixes:
                fix_type = fix.get("type", "").lower()
                
                if fix_type == "font":
                    if self._apply_font_fix(slide, fix):
                        applied_count += 1
                elif fix_type == "position":
                    if self._apply_position_fix(slide, fix):
                        applied_count += 1
                elif fix_type == "color":
                    if self._apply_color_fix(slide, fix):
                        applied_count += 1
                elif fix_type == "chart_label":
                    if self._apply_chart_label_fix(slide, fix):
                        applied_count += 1
                else:
                    logger.warning(f"Unknown fix type: {fix_type}")
            
            # Save updated presentation
            prs.save(pptx_path)
            logger.info(f"Applied {applied_count}/{len(fixes)} fixes to slide {slide_index + 1}")
            
            return applied_count > 0
            
        except Exception as e:
            logger.error(f"Failed to apply fixes: {e}")
            return False
    
    def _apply_font_fix(self, slide, fix: Dict) -> bool:
        """Apply font changes to a shape."""
        try:
            shape_index = fix.get("shape", 0)
            
            if shape_index >= len(slide.shapes):
                return False
            
            shape = slide.shapes[shape_index]
            
            if not shape.has_text_frame:
                return False
            
            for paragraph in shape.text_frame.paragraphs:
                if "size" in fix:
                    paragraph.font.size = Pt(fix["size"])
                if "bold" in fix:
                    paragraph.font.bold = fix["bold"]
                if "italic" in fix:
                    paragraph.font.italic = fix["italic"]
            
            logger.debug(f"Applied font fix to shape {shape_index}")
            return True
            
        except Exception as e:
            logger.debug(f"Font fix failed: {e}")
            return False
    
    def _apply_position_fix(self, slide, fix: Dict) -> bool:
        """Apply position/size changes to a shape."""
        try:
            shape_index = fix.get("shape", 0)
            
            if shape_index >= len(slide.shapes):
                return False
            
            shape = slide.shapes[shape_index]
            
            if "x" in fix:
                shape.left = Inches(fix["x"])
            if "y" in fix:
                shape.top = Inches(fix["y"])
            if "width" in fix:
                shape.width = Inches(fix["width"])
            if "height" in fix:
                shape.height = Inches(fix["height"])
            
            logger.debug(f"Applied position fix to shape {shape_index}")
            return True
            
        except Exception as e:
            logger.debug(f"Position fix failed: {e}")
            return False
    
    def _apply_color_fix(self, slide, fix: Dict) -> bool:
        """Apply color changes to a shape's text."""
        try:
            shape_index = fix.get("shape", 0)
            hex_color = fix.get("color", "#333333").lstrip('#')
            
            if shape_index >= len(slide.shapes):
                return False
            
            shape = slide.shapes[shape_index]
            
            if not shape.has_text_frame:
                return False
            
            rgb = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
            
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.color.rgb = RGBColor(*rgb)
            
            logger.debug(f"Applied color fix to shape {shape_index}")
            return True
            
        except Exception as e:
            logger.debug(f"Color fix failed: {e}")
            return False
    
    def _apply_chart_label_fix(self, slide, fix: Dict) -> bool:
        """Apply font size to chart data labels."""
        try:
            chart_index = fix.get("chart", 0)
            label_size = fix.get("size", 9)
            
            chart_count = 0
            for shape in slide.shapes:
                if shape.has_chart:
                    if chart_count == chart_index:
                        chart = shape.chart
                        plot = chart.plots[0]
                        
                        if plot.has_data_labels:
                            plot.data_labels.font.size = Pt(label_size)
                            logger.debug(f"Applied chart label fix to chart {chart_index}")
                            return True
                    chart_count += 1
            
            return False
            
        except Exception as e:
            logger.debug(f"Chart label fix failed: {e}")
            return False
