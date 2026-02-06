"""
deckAIn - Compliance Validator
Stage 7: Automated compliance checks
"""

from pathlib import Path
from pptx import Presentation

from extractors.schemas import ComplianceReport, FinancialData
from config.settings import ENABLE_COMPLIANCE_CHECKS
from config.sector_config import validate_sector_metrics
from utils.logger import setup_logger, log_stage, log_compliance_check

logger = setup_logger(__name__)

class ComplianceValidator:
    """
    Validate presentation compliance with Kelp requirements.
    """
    
    def __init__(self):
        """Initialize validator."""
        logger.info("ComplianceValidator initialized")
    
    def validate_presentation(
        self,
        pptx_path: str,
        citations_path: str = None,
        company_name: str = None,
        sector: str = None,
        financial_data: FinancialData = None
    ) -> ComplianceReport:
        """
        Run all compliance checks.
        
        Args:
            pptx_path: Path to PPTX file
            citations_path: Path to citations document
            company_name: Original company name (to check for leaks)
        
        Returns:
            Compliance report
        """
        log_stage(logger, 7, "Compliance Validation", "START")
        
        if not ENABLE_COMPLIANCE_CHECKS:
            logger.info(" Compliance checks disabled in config")
            log_stage(logger, 7, "Compliance Validation", "SKIP")
            return None
        
        report = ComplianceReport()
        
        pptx_path = Path(pptx_path)
        if not pptx_path.exists():
            report.add_check(
                "File Exists",
                False,
                f"PPTX not found: {pptx_path}",
                "critical"
            )
            report.calculate_score()
            return report
        
        # Load presentation
        prs = Presentation(str(pptx_path))
        
        # Check 1: Has native charts (CRITICAL)
        has_charts = self._check_native_charts(prs)
        report.add_check(
            "Native Editable Charts",
            has_charts,
            "Found editable charts" if has_charts else "No native charts found",
            "critical"
        )
        log_compliance_check(logger, "Native Charts", has_charts)
        
        # Check 2: Anonymization (CRITICAL)
        is_anonymous = self._check_anonymization(prs, company_name)
        report.add_check(
            "Anonymization",
            is_anonymous,
            "No company name found" if is_anonymous else f"Company name '{company_name}' detected",
            "critical"
        )
        log_compliance_check(logger, "Anonymization", is_anonymous)
        
        # Check 2.5: Sector-specific metrics (IMPORTANT for adaptability)
        if sector and financial_data:
            available_metrics = set()
            if financial_data.revenue_by_year:
                available_metrics.add("revenue")
            if financial_data.ebitda_by_year:
                available_metrics.add("ebitda")
            if financial_data.pat_by_year:
                available_metrics.add("pat")
            if financial_data.facility_count:
                available_metrics.add("facility_count")
            if financial_data.employee_count:
                available_metrics.add("employee_count")
            if financial_data.export_revenue_pct:
                available_metrics.add("export_revenue_pct")
            if financial_data.additional_metrics:
                available_metrics.update(financial_data.additional_metrics.keys())
            
            sector_validation = validate_sector_metrics(sector, available_metrics)
            is_sector_compliant = sector_validation["is_compliant"]
            coverage = sector_validation["coverage_pct"]
            
            report.add_check(
                "Sector-Specific Metrics",
                is_sector_compliant,
                f"{sector} metrics coverage: {coverage:.0f}% ({len(sector_validation['present_kpis'])}/{len(sector_validation['required_kpis'])})",
                "warning"
            )
            log_compliance_check(logger, f"{sector} Metrics", is_sector_compliant)
            
            if not is_sector_compliant:
                logger.warning(f"Missing {sector}-specific KPIs: {', '.join(sector_validation['missing_kpis'])}")
        
        # Check 3: Citations exist
        if citations_path:
            citations_exist = Path(citations_path).exists()
            report.add_check(
                "Citations Document",
                citations_exist,
                "Citations file exists" if citations_exist else "No citations file",
                "warning"
            )
            log_compliance_check(logger, "Citations", citations_exist)
        
        # Check 4: Slide count (per Premium Overhaul: 5 slides)
        slide_count = len(prs.slides)
        correct_count = slide_count == 5
        report.add_check(
            "Slide Count",
            correct_count,
            f"{slide_count} slides" + (" (correct)" if correct_count else " (expected 5)"),
            "info"
        )
        log_compliance_check(logger, "Slide Count", correct_count, f"({slide_count} slides)")
        
        # Check 5: Branding
        has_kelp = self._check_branding(prs)
        report.add_check(
            "Kelp Branding",
            has_kelp,
            "Kelp branding present" if has_kelp else "Kelp branding missing",
            "warning"
        )
        log_compliance_check(logger, "Branding", has_kelp)
        
        # Calculate final score
        report.calculate_score()
        
        log_stage(logger, 7, "Compliance Validation", "DONE")
        logger.info(f"Compliance Score: {report.overall_score:.1f}%")
        logger.info(f"{' PASSED' if report.passed else ' FAILED'}")
        
        return report
    
    def _check_native_charts(self, prs: Presentation) -> bool:
        """Check if presentation has native charts (not images)."""
        
        chart_count = 0
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    chart_count += 1
        
        logger.debug(f"Found {chart_count} native charts")
        return chart_count > 0
    
    def _check_anonymization(self, prs: Presentation, company_name: str = None) -> bool:
        """Check if company name appears in slides."""
        
        if not company_name:
            return True # Can't check without knowing the name
        
        company_lower = company_name.lower()
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.lower()
                    
                    if company_lower in text:
                        logger.warning(f" Company name found in slide text: '{company_name}'")
                        return False
        
        return True
    
    def _check_branding(self, prs: Presentation) -> bool:
        """Check if Kelp branding is present."""
        
        # Check for "Kelp" or "Confidential" in footer/text
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.lower()
                    
                    if "kelp" in text or "confidential" in text:
                        return True
        
        return False
